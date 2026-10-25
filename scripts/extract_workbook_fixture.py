"""Builds tests/fixtures/workbook.json from the reference pricing workbook.

Usage: python3 -I scripts/extract_workbook_fixture.py <workbook.xlsx> <out.json>
Reads cached values (the workbook must have been saved by Excel). Product names are
replaced by neutral labels so no brand data enters the repository.
"""
import json
import sys

import openpyxl

path, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(path, data_only=True)
S, T, P = wb["Settings"], wb["Price Tiers"], wb["Products"]
yes = lambda v: str(v).strip().lower() == "yes"
n = lambda v: float(v) if isinstance(v, (int, float)) else 0.0

fixed = []
for r in range(25, 42):
    fixed.append({"id": f"fc{r}", "name": f"Fixed cost {r - 24}", "category": str(S.cell(r, 2).value or ""),
                  "monthly": n(S.cell(r, 3).value), "perOrder": n(S.cell(r, 4).value)})
markets = []
for r in range(48, 58):
    name = S.cell(r, 1).value
    if not name:
        continue
    markets.append({"name": name, "taxRate": n(S.cell(r, 2).value), "pricesIncludeTax": yes(S.cell(r, 3).value),
                    "dutyRate": n(S.cell(r, 4).value), "otherTaxPct": n(S.cell(r, 5).value),
                    "region": "domestic" if S.cell(r, 6).value == "Domestic" else "international",
                    "customerFxFeePct": n(S.cell(r, 7).value), "shippingCharged": n(S.cell(r, 8).value)})
avail = {"All": "all", "Domestic only": "domestic", "International only": "international"}
methods = []
for r in range(62, 77):
    name = S.cell(r, 1).value
    if not name:
        continue
    methods.append({"name": name, "includeInWorstCase": yes(S.cell(r, 2).value),
                    "availability": avail[S.cell(r, 3).value], "domesticPct": n(S.cell(r, 4).value),
                    "domesticFixed": n(S.cell(r, 5).value), "internationalPct": n(S.cell(r, 6).value),
                    "internationalFixed": n(S.cell(r, 7).value), "extraPct": n(S.cell(r, 8).value)})
tiers = []
for r in range(15, 27):
    if T.cell(r, 2).value is None:
        continue
    tiers.append({"min": n(T.cell(r, 2).value), "max": n(T.cell(r, 3).value),
                  "markupPct": n(T.cell(r, 4).value), "markupAmount": n(T.cell(r, 5).value)})
basis = {"AutoDS: % plus $": "pctPlusAmount", "Higher of % and $": "higherOf",
         "Profit % only": "pctOnly", "Profit $ only": "amountOnly"}[T["B5"].value]
model = {
    "settings": {
        "currency": "USD", "defaultMarket": S["B6"].value, "expectedPaymentMethod": S["B10"].value,
        "fxRate": n(S["B7"].value), "fxFeePct": n(S["B8"].value), "sellerPaysDuty": yes(S["B9"].value),
        "chargebackPct": n(S["B11"].value), "platformFeePct": 0, "platformFeePerOrder": n(S["B12"].value),
        "discountPct": n(S["B13"].value), "otherCostPerOrder": n(S["B14"].value), "refundPct": n(S["B15"].value),
        "returnHandlingCost": n(S["B16"].value), "targetMarginPct": n(S["B17"].value), "bufferPct": n(S["B18"].value),
        "minProfitPerOrder": n(S["B19"].value), "maxBreakEvenRoas": n(S["B20"].value),
        "monthlyOrders": n(S["B21"].value), "paymentFeeBase": "taxInclusive"},
    "pricing": {"tiers": tiers, "basis": basis, "feeGrossUpPct": n(T["B6"].value), "feeGrossUpFixed": n(T["B7"].value),
                "includeShippingInCost": yes(T["B9"].value), "minMarkupProfit": n(T["B8"].value),
                "rounding": "99" if yes(T["B10"].value) else "cent", "raiseToMinimumSafe": yes(T["B11"].value),
                "competitorBandPct": n(T["B12"].value)},
    "fixedCosts": fixed, "markets": markets, "paymentMethods": methods,
}
cols = {"tier": "G", "markupPrice": "H", "markupProfit": "I", "minimumSafePrice": "J", "priceToPublish": "K",
        "competitorDiffPct": "P", "netProfit": "R", "netMargin": "S", "netProfitWorst": "T", "worstMethod": "U",
        "grossProfit": "V", "grossMargin": "W", "breakEvenPrice": "X", "recommendedPrice": "Y", "breakEvenRoas": "Z",
        "maxAdCostAtTarget": "AA", "status": "AB", "netRevenue": "BR", "paymentFeeExpected": "BT",
        "paymentFeeWorst": "BU", "roasForTarget": "CC", "contributionPerOrder": "CE", "landedCost": "BD"}
def code(text):
    """Status style text ("TOO HIGH: above competitor") to its code ("TOO_HIGH")."""
    return str(text).split(":")[0].strip().replace(" ", "_") if text not in (None, "") else None


def price_source(text):
    """Products!CF (internal price source text) to the engine's PriceSource code."""
    if text in (None, ""):
        return None
    t = str(text)
    if t.startswith("Your override"):
        return "override"
    if t.startswith("Raised to minimum safe"):
        return "raisedToMinimumSafe"
    return "markup"


products = []
for r in range(6, 206):
    if P[f"A{r}"].value in (None, ""):
        continue
    i = len(products) + 1
    expected = {k: P[f"{c}{r}"].value for k, c in cols.items()}
    expected["competitorPosition"] = code(P[f"Q{r}"].value)
    expected["priceSource"] = price_source(P[f"CF{r}"].value)
    products.append({
        "input": {"id": f"wb{i}", "name": f"Workbook product {i}", "market": P[f"B{r}"].value or "",
                  "unitCost": P[f"C{r}"].value, "supplierShipping": P[f"D{r}"].value,
                  "adCostPerOrder": P[f"E{r}"].value, "priceOverride": P[f"F{r}"].value,
                  "competitorPrice": P[f"L{r}"].value},
        "expected": expected,
    })
summary = {"fixedMonthly": S["C42"].value, "fixedPerOrder": S["C43"].value}
json.dump({"model": model, "products": products, "summary": summary}, open(out, "w"), indent=1)
print(len(products), "products")
